#!/usr/bin/env python3
"""
extract-pptx.py — Pull every piece of content out of an existing .pptx.

Used in Phase 4 (restyling an existing deck): the extracted JSON is the source of
truth for text, order, images, tables, and speaker notes when the deck is rebuilt
in a new style.

Usage:
    python scripts/extract-pptx.py <input.pptx> [output_dir]

Writes:
    <output_dir>/extracted-slides.json
    <output_dir>/assets/slideN_imgM.<ext>     (every picture, deduplicated by content)

Requires: pip install python-pptx
"""

import hashlib
import json
import os
import sys

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Emu
except ImportError:
    print("python-pptx is required: pip install python-pptx")
    sys.exit(1)

EMU_PER_IN = 914400


def inches(v):
    return round(Emu(v).inches, 3) if v is not None else None


def bbox(shape):
    return {
        "x": inches(shape.left), "y": inches(shape.top),
        "w": inches(shape.width), "h": inches(shape.height),
    }


def paragraphs_of(text_frame):
    paras = []
    for p in text_frame.paragraphs:
        text = "".join(r.text for r in p.runs)
        if not text.strip():
            continue
        paras.append({"text": text, "level": p.level})
    return paras


def walk_shapes(shapes, slide_num, slide_data, assets_dir, seen_hashes, depth=0):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            walk_shapes(shape.shapes, slide_num, slide_data, assets_dir, seen_hashes, depth + 1)
            continue

        entry = {"name": shape.name, "bbox": bbox(shape)}

        if shape.has_text_frame and shape.text_frame.text.strip():
            paras = paragraphs_of(shape.text_frame)
            is_title = False
            try:
                is_title = shape.is_placeholder and shape.placeholder_format.type is not None and \
                    "TITLE" in str(shape.placeholder_format.type)
            except (AttributeError, ValueError):
                pass
            if is_title and not slide_data["title"]:
                slide_data["title"] = " ".join(p["text"] for p in paras)
            else:
                entry.update({"type": "text", "paragraphs": paras})
                slide_data["content"].append(entry)

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(shape, "image", None) is not None:
            try:
                image = shape.image
            except Exception:  # noqa: BLE001 — SVG/EMF or linked pictures
                slide_data["content"].append({**entry, "type": "image", "path": None, "note": "unreadable image (vector or linked)"})
                continue
            digest = hashlib.sha1(image.blob).hexdigest()
            if digest in seen_hashes:
                rel_path = seen_hashes[digest]
            else:
                image_name = f"slide{slide_num}_img{len(slide_data['images']) + 1}.{image.ext}"
                with open(os.path.join(assets_dir, image_name), "wb") as f:
                    f.write(image.blob)
                rel_path = f"assets/{image_name}"
                seen_hashes[digest] = rel_path
            img = {**entry, "type": "image", "path": rel_path}
            slide_data["images"].append(img)
            slide_data["content"].append(img)

        if shape.has_table:
            rows = []
            for r in shape.table.rows:
                rows.append([c.text for c in r.cells])
            slide_data["content"].append({**entry, "type": "table", "rows": rows})

        if shape.has_chart:
            chart = shape.chart
            series = []
            try:
                categories = [str(c) for c in chart.plots[0].categories]
            except Exception:  # noqa: BLE001
                categories = []
            for plot in chart.plots:
                for s in plot.series:
                    try:
                        series.append({"name": s.name, "values": list(s.values)})
                    except Exception:  # noqa: BLE001
                        series.append({"name": s.name, "values": []})
            slide_data["content"].append({
                **entry, "type": "chart",
                "chart_type": str(chart.chart_type), "categories": categories, "series": series,
            })


def extract_pptx(file_path, output_dir="."):
    prs = Presentation(file_path)
    assets_dir = os.path.join(output_dir, "assets")
    os.makedirs(assets_dir, exist_ok=True)
    seen_hashes = {}

    deck = {
        "source": os.path.abspath(file_path),
        "slide_size_in": [inches(prs.slide_width), inches(prs.slide_height)],
        "slides": [],
    }

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_data = {
            "number": slide_num,
            "layout": slide.slide_layout.name,
            "title": "",
            "content": [],
            "images": [],
            "notes": "",
        }
        walk_shapes(slide.shapes, slide_num, slide_data, assets_dir, seen_hashes)
        if slide.has_notes_slide:
            slide_data["notes"] = slide.notes_slide.notes_text_frame.text
        deck["slides"].append(slide_data)

    return deck


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python extract-pptx.py <input.pptx> [output_dir]")
        sys.exit(1)

    input_file = sys.argv[1]
    output_dir = sys.argv[2] if len(sys.argv) > 2 else "."
    os.makedirs(output_dir, exist_ok=True)

    deck = extract_pptx(input_file, output_dir)

    output_path = os.path.join(output_dir, "extracted-slides.json")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(deck, f, indent=2, ensure_ascii=False)

    print(f"Extracted {len(deck['slides'])} slides to {output_path}")
    print(f"Slide size: {deck['slide_size_in'][0]} x {deck['slide_size_in'][1]} in")
    for s in deck["slides"]:
        kinds = {}
        for c in s["content"]:
            kinds[c["type"]] = kinds.get(c["type"], 0) + 1
        summary = ", ".join(f"{v} {k}" for k, v in kinds.items()) or "empty"
        print(f"  Slide {s['number']}: {s['title'] or '(no title)'} — {summary}")
